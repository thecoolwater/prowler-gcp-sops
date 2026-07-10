"""
Generate six separate simple, plain DOCX SOP documents:
1. Customer Prereqs: AWS (no Prowler references, CloudFormation screenshots)
2. Customer Prereqs: GCP (no Prowler references, GCP console screenshots)
3. Customer Prereqs: Azure (no Prowler references, Azure console screenshots)
4. Internal Audit: AWS (retains Prowler, AWS CLI, Prowler App setup)
5. Internal Audit: GCP (retains Prowler, GCP CLI, Prowler App setup)
6. Internal Audit: Azure (retains Prowler, Azure CLI, Prowler App setup)
"""

from docx import Document
from docx.shared import Inches, Pt
import datetime

TODAY = datetime.date.today().strftime("%B %d, %Y")

def add_simple_header(doc, title, doc_id):
    """Add a simple title and document ID at the top of the first page."""
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.name = "Arial"
    
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"Document ID: {doc_id}  |  Date: {TODAY}  |  Version: 1.0\n")
    r_meta.font.size = Pt(9.5)
    r_meta.italic = True
    r_meta.font.name = "Arial"

def add_simple_heading(doc, text, level):
    """Add plain Arial headings."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Arial"
        run.bold = True
        if level == 1:
            run.font.size = Pt(14)
        elif level == 2:
            run.font.size = Pt(12)
        else:
            run.font.size = Pt(11)

def add_simple_paragraph(doc, text, bold=False, italic=False, space_after=6):
    """Add simple Arial body text."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.bold = bold
    run.italic = italic
    return p

def add_simple_bullet(doc, text):
    """Add plain bullet points."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.clear()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)

def add_simple_code(doc, code_text):
    """Add a plain, unshaded code block."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)

def add_simple_table(doc, headers, rows):
    """Add a simple unstyled table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True
    
    # Header
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(header)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        
    # Rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(str(cell_text))
            run.font.name = "Arial"
            run.font.size = Pt(9)
            
    doc.add_paragraph() # space after table

def add_simple_image(doc, img_path, caption):
    """Add an image with a simple caption centered."""
    try:
        p_img = doc.add_paragraph()
        p_img.alignment = 1 # Center
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=Inches(5.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = 1 # Center
        p_cap.paragraph_format.space_after = Pt(8)
        run = p_cap.add_run(f"Figure: {caption}")
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.italic = True
    except Exception as e:
        print(f"Warning: Could not add image {img_path} - {e}")

# ═══════════════════════════════════════════════════════════════════════════════
# GENERATE CUSTOMER-FACING DOCUMENTS
# ═══════════════════════════════════════════════════════════════════════════════

def make_aws_customer_sop():
    doc = Document()
    DOC_ID = "SOP-AWS-PREREQUISITES-CUSTOMER"
    add_simple_header(doc, "AWS Prerequisites Setup for Cloud Security Assessment", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This document outlines the steps required to configure read-only access for a cloud security posture assessment in Amazon Web Services (AWS). All steps are performed on the customer side prior to the assessment.")
    
    add_simple_heading(doc, "2. Setup Access via CloudFormation Role (Recommended)", level=1)
    add_simple_paragraph(doc, "Deploying a CloudFormation stack is the easiest way to generate the required read-only assessment role:")
    
    add_simple_bullet(doc, "Open AWS Console and navigate to CloudFormation > Click 'Create stack' (with new resources).")
    add_simple_image(doc, "cloudformation-nav.png", "CloudFormation Dashboard in AWS Console")
    
    add_simple_bullet(doc, "Select 'Upload a template file' and upload the audit template YAML file.")
    add_simple_image(doc, "upload-template-file.png", "Uploading the CloudFormation template file")
    
    add_simple_bullet(doc, "Enter a stack name (e.g. 'security-audit-role-stack') and enter the target Assessment Account ID if cross-account access is required.")
    add_simple_image(doc, "fill-stack-data.png", "Entering the Stack Parameters")
    
    add_simple_bullet(doc, "Review settings, check 'I acknowledge that AWS CloudFormation might create IAM resources', and click 'Submit'.")
    add_simple_image(doc, "create-stack.png", "Deploying the Stack")
    
    add_simple_bullet(doc, "Once status shows CREATE_COMPLETE, go to Outputs and copy the Role ARN (e.g. arn:aws:iam::<ACCOUNT_ID>:role/security-audit-role).")

    add_simple_heading(doc, "3. Setup Access via IAM User Access Keys (Alternative)", level=1)
    add_simple_paragraph(doc, "If role assumption is not supported, create a local IAM User:")
    add_simple_bullet(doc, "Navigate to IAM > Users > Click 'Create user'.")
    add_simple_bullet(doc, "Name the user 'security-audit-user' and click Next.")
    add_simple_bullet(doc, "Select 'Attach policies directly' > Attach the standard 'SecurityAudit' and 'ViewOnlyAccess' policies.")
    add_simple_bullet(doc, "Click Create user > Select the created user > 'Security credentials' tab > Click 'Create access key'.")
    add_simple_bullet(doc, "Select 'Command Line Interface (CLI)' > Copy the Access Key ID and Secret Access Key.")
    
    add_simple_paragraph(doc, "Deliver the Role ARN OR the Access Key credentials securely to the assessment team.", bold=True)
    
    doc.save("SOP-AWS-PREREQUISITES-CUSTOMER-v1.0.docx")
    print("Generated SOP-AWS-PREREQUISITES-CUSTOMER-v1.0.docx")


def make_gcp_customer_sop():
    doc = Document()
    DOC_ID = "SOP-GCP-PREREQUISITES-CUSTOMER"
    add_simple_header(doc, "GCP Prerequisites Setup for Cloud Security Assessment", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This document outlines the steps required to configure read-only access for a cloud security posture assessment in Google Cloud Platform (GCP). All steps are performed on the customer side prior to the assessment.")
    
    add_simple_heading(doc, "2. Enable Required APIs", level=1)
    add_simple_paragraph(doc, "You must enable the Identity and Access Management (IAM) API and Cloud Resource Manager API in your target project:")
    add_simple_code(doc, "gcloud services enable iam.googleapis.com --project <YOUR_PROJECT_ID>\ngcloud services enable cloudresourcemanager.googleapis.com --project <YOUR_PROJECT_ID>")
    
    add_simple_heading(doc, "3. Create a Service Account", level=1)
    add_simple_paragraph(doc, "Create a dedicated service account that the assessment team will use:")
    add_simple_code(doc, "gcloud iam service-accounts create security-audit-sa --display-name=\"Security Audit Service Account\" --project=<YOUR_PROJECT_ID>")
    add_simple_image(doc, "create-service-account.png", "Creating the Service Account in Google Cloud Console")
    
    add_simple_heading(doc, "4. Assign Read-Only Permissions", level=1)
    add_simple_paragraph(doc, "Assign 'Viewer' and 'Service Usage Consumer' roles to the service account:")
    add_simple_code(doc, 
        "gcloud projects add-iam-policy-binding <YOUR_PROJECT_ID> --member=\"serviceAccount:security-audit-sa@<YOUR_PROJECT_ID>.iam.gserviceaccount.com\" --role=\"roles/viewer\"\n\n"
        "gcloud projects add-iam-policy-binding <YOUR_PROJECT_ID> --member=\"serviceAccount:security-audit-sa@<YOUR_PROJECT_ID>.iam.gserviceaccount.com\" --role=\"roles/serviceusage.serviceUsageConsumer\""
    )
    add_simple_image(doc, "service-account-permissions.png", "Assigning Roles in IAM Console")
    
    add_simple_heading(doc, "5. Generate JSON Key Credentials", level=1)
    add_simple_paragraph(doc, "Navigate to the Keys tab under the service account, click 'ADD KEY', choose 'Create new key', and select JSON format:")
    add_simple_image(doc, "create-new-key.png", "Adding a Key to the Service Account")
    add_simple_image(doc, "json-key.png", "Downloading the JSON Key File")
    add_simple_paragraph(doc, "This downloads a credential JSON file. Share this file securely with the assessment team.")
    
    doc.save("SOP-GCP-PREREQUISITES-CUSTOMER-v1.0.docx")
    print("Generated SOP-GCP-PREREQUISITES-CUSTOMER-v1.0.docx")


def make_azure_customer_sop():
    doc = Document()
    DOC_ID = "SOP-AZURE-PREREQUISITES-CUSTOMER"
    add_simple_header(doc, "Azure Prerequisites Setup for Cloud Security Assessment", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This document outlines the steps required to configure read-only access for a cloud security posture assessment in Microsoft Azure. All steps are performed on the customer side prior to the assessment.")
    
    add_simple_heading(doc, "2. App Registration (Microsoft Entra ID)", level=1)
    add_simple_paragraph(doc, "Register a new application in Entra ID (formerly Azure Active Directory) to allow API access:")
    add_simple_bullet(doc, "Navigate to Microsoft Entra ID > App registrations > Click '+ New registration'.")
    add_simple_bullet(doc, "Name it 'security-audit-app' and select 'Accounts in this organizational directory only'.")
    add_simple_bullet(doc, "Copy and save the Application (client) ID and Directory (tenant) ID from the App details page.")
    add_simple_image(doc, "azure-app-registration.jpg", "Azure App Registration details screen showing IDs")
    
    add_simple_heading(doc, "3. Generate Client Secret", level=1)
    add_simple_paragraph(doc, "Generate a client secret key to authenticate the App Registration:")
    add_simple_bullet(doc, "Go to Certificates & secrets > Click '+ New client secret'.")
    add_simple_bullet(doc, "Copy the secret 'Value' immediately (it will be hidden after you navigate away).")
    
    add_simple_heading(doc, "4. Assign Subscription Reader Role", level=1)
    add_simple_paragraph(doc, "Grant the App Registration access to read subscription resource configurations:")
    add_simple_bullet(doc, "Navigate to Subscriptions > Select your subscription > Access control (IAM).")
    add_simple_bullet(doc, "Click 'Add role assignment' > Search for and select the 'Reader' role.")
    add_simple_bullet(doc, "Assign access to 'User, group, or service principal' and select 'security-audit-app'.")
    add_simple_bullet(doc, "Review and click Save to apply the permissions.")
    add_simple_image(doc, "azure-permissions.jpg", "Assigning Subscription Reader Permissions to App Principal")
    
    add_simple_paragraph(doc, "Deliver the Application ID, Tenant ID, and Client Secret safely to the security assessment team.", bold=True)
    
    doc.save("SOP-AZURE-PREREQUISITES-CUSTOMER-v1.0.docx")
    print("Generated SOP-AZURE-PREREQUISITES-CUSTOMER-v1.0.docx")


# ═══════════════════════════════════════════════════════════════════════════════
# GENERATE INTERNAL-ONLY AUDIT EXECUTION DOCUMENTS
# ═══════════════════════════════════════════════════════════════════════════════

def make_aws_internal_sop():
    doc = Document()
    DOC_ID = "SOP-AWS-AUDIT-INTERNAL"
    add_simple_header(doc, "AWS Cloud Security Assessment: Prowler Audit Execution Guide", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This guide details how to execute security scans on AWS using Prowler CLI and deploy the Prowler Web App using Docker Compose.")
    
    add_simple_heading(doc, "2. Running Scans via Prowler CLI", level=1)
    add_simple_paragraph(doc, "Set the customer's shared Access Key credentials as environment variables and run:")
    add_simple_code(doc, 
        "export AWS_ACCESS_KEY_ID=\"<CUSTOMER_ACCESS_KEY>\"\n"
        "export AWS_SECRET_ACCESS_KEY=\"<CUSTOMER_SECRET_KEY>\"\n"
        "prowler aws"
    )
    add_simple_paragraph(doc, "Alternatively, if the customer provided an IAM Role ARN, run:")
    add_simple_code(doc, "prowler aws --role arn:aws:iam::<ACCOUNT_ID>:role/security-audit-role")
    
    add_simple_heading(doc, "3. Setting up the Prowler Web App", level=1)
    add_simple_paragraph(doc, "Download the Docker configuration files:")
    add_simple_code(doc,
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/docker-compose.yml\n"
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/.env"
    )
    add_simple_image(doc, "prowler-env-file.jpg", "Prowler environment config showing port settings and environment options")
    
    add_simple_paragraph(doc, "Spin up the Docker containers in detached mode:")
    add_simple_code(doc, "sudo docker compose up -d")
    add_simple_paragraph(doc, "Open browser to http://localhost:3000. Under 'Cloud Providers', select AWS, enter your keys or Role details:")
    add_simple_image(doc, "prowler-app-dashboard.jpg", "Prowler App main login/providers page")
    
    add_simple_heading(doc, "4. Prowler App Web UI Views", level=1)
    add_simple_paragraph(doc, "Once the scan is complete, view the dashboards and compliance results:")
    add_simple_image(doc, "prowler-dashboard-web.jpg", "Interactive Prowler App Dashboard UI")
    add_simple_image(doc, "compliance-dashboard.jpg", "Compliance dashboard detailing CIS standard posture ratios")
    
    doc.save("SOP-AWS-AUDIT-INTERNAL-v1.0.docx")
    print("Generated SOP-AWS-AUDIT-INTERNAL-v1.0.docx")


def make_gcp_internal_sop():
    doc = Document()
    DOC_ID = "SOP-GCP-AUDIT-INTERNAL"
    add_simple_header(doc, "GCP Cloud Security Assessment: Prowler Audit Execution Guide", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This guide details how to execute security scans on GCP using Prowler CLI and deploy the Prowler Web App using Docker Compose.")
    
    add_simple_heading(doc, "2. Running Scans via Prowler CLI", level=1)
    add_simple_paragraph(doc, "Import the customer's shared 'security-audit-key.json' file and run:")
    add_simple_code(doc, 
        "export GOOGLE_APPLICATION_CREDENTIALS=\"/path/to/security-audit-key.json\"\n"
        "export GOOGLE_CLOUD_QUOTA_PROJECT=\"<CUSTOMER_PROJECT_ID>\"\n"
        "prowler gcp"
    )
    
    add_simple_heading(doc, "3. Setting up the Prowler Web App", level=1)
    add_simple_paragraph(doc, "Download the Docker configuration files:")
    add_simple_code(doc,
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/docker-compose.yml\n"
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/.env"
    )
    add_simple_image(doc, "prowler-env-file.jpg", "Prowler environment config showing port settings and environment options")
    
    add_simple_paragraph(doc, "Spin up the Docker containers in detached mode:")
    add_simple_code(doc, "sudo docker compose up -d")
    add_simple_paragraph(doc, "Open browser to http://localhost:3000. Under 'Cloud Providers', select GCP and upload the JSON credentials:")
    add_simple_image(doc, "prowler-app-dashboard.jpg", "Prowler App main login/providers page")
    
    add_simple_heading(doc, "4. Prowler App Web UI Views", level=1)
    add_simple_paragraph(doc, "Once the scan is complete, view the dashboards and compliance results:")
    add_simple_image(doc, "prowler-dashboard-web.jpg", "Interactive Prowler App Dashboard UI")
    add_simple_image(doc, "compliance-dashboard.jpg", "Compliance dashboard detailing CIS standard posture ratios")
    
    doc.save("SOP-GCP-AUDIT-INTERNAL-v1.0.docx")
    print("Generated SOP-GCP-AUDIT-INTERNAL-v1.0.docx")


def make_azure_internal_sop():
    doc = Document()
    DOC_ID = "SOP-AZURE-AUDIT-INTERNAL"
    add_simple_header(doc, "Azure Cloud Security Assessment: Prowler Audit Execution Guide", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This guide details how to execute security scans on Azure using Prowler CLI and deploy the Prowler Web App using Docker Compose.")
    
    add_simple_heading(doc, "2. Running Scans via Prowler CLI", level=1)
    add_simple_paragraph(doc, "Configure the customer's tenant, client ID, and secret as environment variables and run Prowler with the Service Principal flag:")
    add_simple_code(doc,
        "export AZURE_CLIENT_ID=\"<CUSTOMER_CLIENT_ID>\"\n"
        "export AZURE_TENANT_ID=\"<CUSTOMER_TENANT_ID>\"\n"
        "export AZURE_CLIENT_SECRET=\"<CUSTOMER_CLIENT_SECRET>\"\n"
        "prowler azure --sp-env-auth"
    )
    
    add_simple_heading(doc, "3. Setting up the Prowler Web App", level=1)
    add_simple_paragraph(doc, "Download the Docker configuration files:")
    add_simple_code(doc,
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/docker-compose.yml\n"
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/.env"
    )
    add_simple_image(doc, "prowler-env-file.jpg", "Prowler environment config showing port settings and environment options")
    
    add_simple_paragraph(doc, "Spin up the Docker containers in detached mode:")
    add_simple_code(doc, "sudo docker compose up -d")
    add_simple_paragraph(doc, "Open browser to http://localhost:3000. Under 'Cloud Providers', select Azure and enter the App credentials:")
    add_simple_image(doc, "prowler-app-dashboard.jpg", "Prowler App main login/providers page")
    
    add_simple_heading(doc, "4. Prowler App Web UI Views", level=1)
    add_simple_paragraph(doc, "Once the scan is complete, view the dashboards and compliance results:")
    add_simple_image(doc, "prowler-dashboard-web.jpg", "Interactive Prowler App Dashboard UI")
    add_simple_image(doc, "compliance-dashboard.jpg", "Compliance dashboard detailing CIS standard posture ratios")
    add_simple_image(doc, "azure-misconfigurations.jpg", "Reviewing Azure misconfigurations list in UI")
    
    doc.save("SOP-AZURE-AUDIT-INTERNAL-v1.0.docx")
    print("Generated SOP-AZURE-AUDIT-INTERNAL-v1.0.docx")


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN EXECUTION
# ═══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    make_aws_customer_sop()
    make_gcp_customer_sop()
    make_azure_customer_sop()
    make_aws_internal_sop()
    make_gcp_internal_sop()
    make_azure_internal_sop()
