"""
Generate the AWS & Azure SOP documents matching the Cloud Infrastructure Services structure.
- Document 1: Customer-facing AWS & Azure Prerequisites (No Prowler mentions, embedded AWS & Azure console screenshots)
- Document 2: Internal Prowler Execution Guide (Retains Prowler CLI and Prowler App setup, embedded Prowler App UI screenshots)
"""

from docx import Document
from docx.shared import Inches, Pt
import datetime

TODAY = datetime.date.today().strftime("%B %d, %Y")

def add_simple_header(doc, title, doc_id):
    """Add a simple title and document ID at the top of the first page."""
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.name = "Arial"
    
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"Document ID: {doc_id}  |  Date: {TODAY}  |  Version: 1.0\n")
    r_meta.font.size = Pt(9.5)
    r_meta.italic = True
    r_meta.font.name = "Arial"

def add_simple_heading(doc, text, level):
    """Add plain Arial headings."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Arial"
        run.bold = True
        if level == 1:
            run.font.size = Pt(14)
        elif level == 2:
            run.font.size = Pt(12)
        else:
            run.font.size = Pt(11)

def add_simple_paragraph(doc, text, bold=False, italic=False, space_after=6):
    """Add simple Arial body text."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.bold = bold
    run.italic = italic
    return p

def add_simple_bullet(doc, text):
    """Add plain bullet points."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.clear()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)

def add_simple_code(doc, code_text):
    """Add a plain, unshaded code block."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)

def add_simple_table(doc, headers, rows):
    """Add a simple unstyled table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True
    
    # Header
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(header)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        
    # Rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(str(cell_text))
            run.font.name = "Arial"
            run.font.size = Pt(9)
            
    doc.add_paragraph() # space after table

def add_simple_image(doc, img_path, caption):
    """Add an image with a simple caption centered."""
    try:
        p_img = doc.add_paragraph()
        p_img.alignment = 1 # Center
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=Inches(5.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = 1 # Center
        p_cap.paragraph_format.space_after = Pt(8)
        run = p_cap.add_run(f"Figure: {caption}")
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.italic = True
    except Exception as e:
        print(f"Warning: Could not add image {img_path} - {e}")

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 1: CUSTOMER-FACING AWS & AZURE PREREQUISITES
# ═══════════════════════════════════════════════════════════════════════════════
def make_customer_sop():
    doc = Document()
    DOC_ID = "SOP-AWS-AZURE-SECURITY-PREREQS"
    
    add_simple_header(doc, "AWS & Azure Prerequisites Setup for Cloud Security Assessment", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This document outlines the steps required to configure read-only access for a cloud security posture assessment across Amazon Web Services (AWS) and Microsoft Azure. These configurations are performed on the customer side prior to the assessment.")
    
    # --- AWS PART ---
    add_simple_heading(doc, "2. Amazon Web Services (AWS) Setup Steps", level=1)
    add_simple_paragraph(doc, "To perform the security assessment on AWS, we require read-only access to your AWS metadata. You can configure this either by deploying an IAM Role (Recommended) or creating an IAM User with Access Keys.")
    
    add_simple_heading(doc, "2.1 Method A: Create IAM Role via CloudFormation (Recommended)", level=2)
    add_simple_paragraph(doc, "Using a CloudFormation template is the fastest way to create the read-only assessment role:")
    add_simple_bullet(doc, "Open AWS Console and navigate to CloudFormation > Click 'Create stack' (with new resources).")
    add_simple_image(doc, "cloudformation-nav.png", "CloudFormation Dashboard in AWS Console")
    
    add_simple_bullet(doc, "Select 'Upload a template file' and upload the assessment role template YAML file.")
    add_simple_image(doc, "upload-template-file.png", "Uploading the CloudFormation template file")
    
    add_simple_bullet(doc, "Specify a stack name (e.g. 'security-audit-role-stack') and enter the target Assessment Account ID if doing cross-account role assumption.")
    add_simple_image(doc, "fill-stack-data.png", "Entering the Stack Parameters")
    
    add_simple_bullet(doc, "Review the settings, check the box 'I acknowledge that AWS CloudFormation might create IAM resources', and click 'Submit' (or Create stack).")
    add_simple_image(doc, "create-stack.png", "Deploying the Stack")
    
    add_simple_bullet(doc, "Once the stack status is CREATE_COMPLETE, go to Outputs and copy the Role ARN (e.g. arn:aws:iam::<ACCOUNT_ID>:role/security-audit-role).")

    add_simple_heading(doc, "2.2 Method B: Create IAM User with Access Keys (Alternative)", level=2)
    add_simple_paragraph(doc, "If role assumption is not possible, create a local IAM User:")
    add_simple_bullet(doc, "Navigate to IAM > Users > Click 'Create user'.")
    add_simple_bullet(doc, "Name the user 'security-audit-user' and click Next.")
    add_simple_bullet(doc, "Select 'Attach policies directly' > Search for and attach the standard 'SecurityAudit' and 'ViewOnlyAccess' policies.")
    add_simple_bullet(doc, "Review and click 'Create user'.")
    add_simple_bullet(doc, "Select the created user > Go to 'Security credentials' tab > Click 'Create access key'.")
    add_simple_bullet(doc, "Select 'Command Line Interface (CLI)' > Copy the Access Key ID and Secret Access Key.")
    
    add_simple_paragraph(doc, "Deliver the Role ARN (Method A) OR the Access Key ID and Secret Access Key (Method B) securely to the assessment team.", bold=True)

    # --- AZURE PART ---
    add_simple_heading(doc, "3. Microsoft Azure Setup Steps", level=1)
    
    add_simple_heading(doc, "3.1 App Registration (Microsoft Entra ID)", level=2)
    add_simple_paragraph(doc, "Register a new application in Entra ID (formerly Azure Active Directory) to allow API access:")
    add_simple_bullet(doc, "Navigate to Microsoft Entra ID > App registrations > Click '+ New registration'.")
    add_simple_bullet(doc, "Name it 'security-audit-app' and select 'Accounts in this organizational directory only'.")
    add_simple_bullet(doc, "Copy and save the Application (client) ID and Directory (tenant) ID from the App details page.")
    add_simple_image(doc, "azure-app-registration.jpg", "Azure App Registration details screen showing IDs")
    
    add_simple_heading(doc, "3.2 Generate Client Secret", level=2)
    add_simple_paragraph(doc, "Generate a client secret key to authenticate the App Registration:")
    add_simple_bullet(doc, "Go to Certificates & secrets > Click '+ New client secret'.")
    add_simple_bullet(doc, "Copy the secret 'Value' immediately (it will be hidden after you navigate away).")
    
    add_simple_heading(doc, "3.3 Assign Subscription Reader Role", level=2)
    add_simple_paragraph(doc, "Grant the App Registration access to read subscription resource configurations:")
    add_simple_bullet(doc, "Navigate to Subscriptions > Select your subscription > Access control (IAM).")
    add_simple_bullet(doc, "Click 'Add role assignment' > Search for and select the 'Reader' role.")
    add_simple_bullet(doc, "Assign access to 'User, group, or service principal' and select 'security-audit-app'.")
    add_simple_bullet(doc, "Review and click Save to apply the permissions.")
    add_simple_image(doc, "azure-permissions.jpg", "Assigning Subscription Reader Permissions to App Principal")
    
    add_simple_paragraph(doc, "Deliver the Application ID, Tenant ID, and Client Secret safely to the security assessment team.", bold=True)
    
    doc.save("SOP-AWS-AZURE-PREREQUISITES-CUSTOMER-v1.0.docx")
    print("Generated SOP-AWS-AZURE-PREREQUISITES-CUSTOMER-v1.0.docx")

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 2: INTERNAL PROWLER SCAN EXECUTION & APP DEPLOYMENT
# ═══════════════════════════════════════════════════════════════════════════════
def make_internal_sop():
    doc = Document()
    DOC_ID = "SOP-Prowler-AWS-AZURE-AUDIT-INTERNAL"
    
    add_simple_header(doc, "Cloud Security Assessment: Prowler AWS & Azure Audit Execution Guide", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This internal guide details how to execute security assessment scans using Prowler CLI and how to set up the interactive self-hosted Prowler App using Docker Compose.")
    
    # --- Prowler CLI ---
    add_simple_heading(doc, "2. Running Scans via Prowler CLI", level=1)
    
    add_simple_heading(doc, "2.1 AWS Scan Execution", level=2)
    add_simple_paragraph(doc, "To run the AWS audit using Access Keys, set environment variables and run:")
    add_simple_code(doc, 
        "export AWS_ACCESS_KEY_ID=\"<CUSTOMER_ACCESS_KEY>\"\n"
        "export AWS_SECRET_ACCESS_KEY=\"<CUSTOMER_SECRET_KEY>\"\n"
        "prowler aws"
    )
    add_simple_paragraph(doc, "Alternatively, if the customer provided an IAM Role ARN for cross-account scans, run:")
    add_simple_code(doc, "prowler aws --role arn:aws:iam::<ACCOUNT_ID>:role/security-audit-role")
    
    add_simple_heading(doc, "2.2 Azure Scan Execution", level=2)
    add_simple_paragraph(doc, "To run the Azure audit, configure the customer's tenant, client ID, and secret as environment variables and run Prowler with the Service Principal flag:")
    add_simple_code(doc,
        "export AZURE_CLIENT_ID=\"<CUSTOMER_CLIENT_ID>\"\n"
        "export AZURE_TENANT_ID=\"<CUSTOMER_TENANT_ID>\"\n"
        "export AZURE_CLIENT_SECRET=\"<CUSTOMER_CLIENT_SECRET>\"\n"
        "prowler azure --sp-env-auth"
    )

    # --- Prowler App Docker ---
    add_simple_heading(doc, "3. Setting up the Prowler Web App", level=1)
    add_simple_paragraph(doc, "The Prowler App provides a clean, collaborative dashboard web UI for multi-cloud audits.")
    
    add_simple_heading(doc, "3.1 Download Docker Compose Files", level=2)
    add_simple_paragraph(doc, "Use curl to download the official Prowler Docker configuration files:")
    add_simple_code(doc,
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/docker-compose.yml\n"
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/.env"
    )
    
    add_simple_heading(doc, "3.2 Configure Environment (.env)", level=2)
    add_simple_paragraph(doc, "Open the '.env' file to configure system settings, users, and passwords:")
    add_simple_image(doc, "prowler-env-file.jpg", "Prowler environment config showing port settings and environment options")
    
    add_simple_heading(doc, "3.3 Run the Web App", level=2)
    add_simple_paragraph(doc, "Spin up the Docker containers in detached mode:")
    add_simple_code(doc, "sudo docker compose up -d")
    add_simple_paragraph(doc, "Access the web panel in your browser at http://localhost:3000 (or http://SERVER_IP:3000). Create an admin login upon initial startup.")
    
    # --- Prowler Web App Usage ---
    add_simple_heading(doc, "4. Prowler App Web UI Usage", level=1)
    
    add_simple_heading(doc, "4.1 Connecting Cloud Accounts", level=2)
    add_simple_paragraph(doc, "Log into the Prowler App, navigate to Cloud Providers, select AWS or Azure, and enter the client keys or IAM Role details:")
    add_simple_image(doc, "prowler-app-dashboard.jpg", "Prowler App main login/providers page")
    
    add_simple_heading(doc, "4.2 Performing Assessments and Dashboard Views", level=2)
    add_simple_paragraph(doc, "Run the assessment job. The Prowler Dashboard will populate showing Findings by Severity and findings by provider:")
    add_simple_image(doc, "prowler-dashboard-web.jpg", "Interactive Prowler App Dashboard UI")
    
    add_simple_heading(doc, "4.3 Compliance Mappings", level=2)
    add_simple_paragraph(doc, "Navigate to Compliance > Select CIS Benchmark, NIST, or ISO 27001 to view compliance ratios:")
    add_simple_image(doc, "compliance-dashboard.jpg", "Compliance dashboard detailing CIS standard posture ratios")
    
    add_simple_heading(doc, "4.4 Reviewing Misconfigurations", level=2)
    add_simple_paragraph(doc, "Review the active failures listing directly from the UI and download reports:")
    add_simple_image(doc, "azure-misconfigurations.jpg", "Misconfigurations page showing status, severity, and mitigation instructions")
    
    doc.save("SOP-AWS-AZURE-AUDIT-INTERNAL-v1.0.docx")
    print("Generated SOP-AWS-AZURE-AUDIT-INTERNAL-v1.0.docx")

if __name__ == "__main__":
    make_customer_sop()
    make_internal_sop()
