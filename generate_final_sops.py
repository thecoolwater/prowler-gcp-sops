"""
Generate the final GCP/Azure SOP documents matching the Cloud Infrastructure Services blog post structure.
- Document 1: Customer-facing GCP & Azure Prerequisites (No Prowler mentions, embedded GCP & Azure console screenshots)
- Document 2: Internal Prowler Execution Guide (Retains Prowler CLI and Prowler App setup, embedded Prowler App UI screenshots)
"""

from docx import Document
from docx.shared import Inches, Pt
import datetime

TODAY = datetime.date.today().strftime("%B %d, %Y")

def add_simple_header(doc, title, doc_id):
    """Add a simple title and document ID at the top of the first page."""
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.name = "Arial"
    
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"Document ID: {doc_id}  |  Date: {TODAY}  |  Version: 1.0\n")
    r_meta.font.size = Pt(9.5)
    r_meta.italic = True
    r_meta.font.name = "Arial"

def add_simple_heading(doc, text, level):
    """Add plain Arial headings."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Arial"
        run.bold = True
        if level == 1:
            run.font.size = Pt(14)
        elif level == 2:
            run.font.size = Pt(12)
        else:
            run.font.size = Pt(11)

def add_simple_paragraph(doc, text, bold=False, italic=False, space_after=6):
    """Add simple Arial body text."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.bold = bold
    run.italic = italic
    return p

def add_simple_bullet(doc, text):
    """Add plain bullet points."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.clear()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)

def add_simple_code(doc, code_text):
    """Add a plain, unshaded code block."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)

def add_simple_table(doc, headers, rows):
    """Add a simple unstyled table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True
    
    # Header
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(header)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        
    # Rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(str(cell_text))
            run.font.name = "Arial"
            run.font.size = Pt(9)
            
    doc.add_paragraph() # space after table

def add_simple_image(doc, img_path, caption):
    """Add an image with a simple caption centered."""
    try:
        p_img = doc.add_paragraph()
        p_img.alignment = 1 # Center
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=Inches(5.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = 1 # Center
        p_cap.paragraph_format.space_after = Pt(8)
        run = p_cap.add_run(f"Figure: {caption}")
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.italic = True
    except Exception as e:
        print(f"Warning: Could not add image {img_path} - {e}")

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 1: CUSTOMER-FACING GCP & AZURE PREREQUISITES
# ═══════════════════════════════════════════════════════════════════════════════
def make_customer_sop():
    doc = Document()
    DOC_ID = "SOP-CLOUD-SECURITY-PREREQS-CUSTOMER"
    
    add_simple_header(doc, "GCP & Azure Prerequisites Setup for Cloud Security Assessment", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This document outlines the steps required to configure read-only access for a cloud security assessment across Google Cloud Platform (GCP) and Microsoft Azure. These configurations are performed on the customer side prior to the assessment.")
    
    # --- GCP PART ---
    add_simple_heading(doc, "2. Google Cloud Platform (GCP) Setup Steps", level=1)
    
    add_simple_heading(doc, "2.1 Enable Required APIs", level=2)
    add_simple_paragraph(doc, "You must enable the Identity and Access Management (IAM) API and Cloud Resource Manager API in your target project. You can enable them via Cloud Console or using these commands:")
    add_simple_code(doc, "gcloud services enable iam.googleapis.com --project <YOUR_PROJECT_ID>\ngcloud services enable cloudresourcemanager.googleapis.com --project <YOUR_PROJECT_ID>")
    
    add_simple_heading(doc, "2.2 Create Service Account", level=2)
    add_simple_paragraph(doc, "Create a dedicated service account that the assessment team will use:")
    add_simple_code(doc, "gcloud iam service-accounts create security-audit-sa --display-name=\"Security Audit Service Account\" --project=<YOUR_PROJECT_ID>")
    add_simple_image(doc, "create-service-account.png", "Creating the Service Account in Google Cloud Console")
    
    add_simple_heading(doc, "2.3 Assign Read-Only Permissions", level=2)
    add_simple_paragraph(doc, "Assign 'Viewer' and 'Service Usage Consumer' roles to the service account:")
    add_simple_code(doc, 
        "gcloud projects add-iam-policy-binding <YOUR_PROJECT_ID> --member=\"serviceAccount:security-audit-sa@<YOUR_PROJECT_ID>.iam.gserviceaccount.com\" --role=\"roles/viewer\"\n\n"
        "gcloud projects add-iam-policy-binding <YOUR_PROJECT_ID> --member=\"serviceAccount:security-audit-sa@<YOUR_PROJECT_ID>.iam.gserviceaccount.com\" --role=\"roles/serviceusage.serviceUsageConsumer\""
    )
    add_simple_image(doc, "service-account-permissions.png", "Assigning Roles in IAM Console")
    
    add_simple_heading(doc, "2.4 Generate JSON Key Credentials", level=2)
    add_simple_paragraph(doc, "Navigate to the Keys tab under the service account, click 'ADD KEY', choose 'Create new key', and select JSON format:")
    add_simple_image(doc, "create-new-key.png", "Adding a Key to the Service Account")
    add_simple_image(doc, "json-key.png", "Downloading the JSON Key File")
    add_simple_paragraph(doc, "This downloads a credential JSON file. Share this file securely with the assessment team.")
    
    # --- AZURE PART ---
    add_simple_heading(doc, "3. Microsoft Azure Setup Steps", level=1)
    
    add_simple_heading(doc, "3.1 App Registration (Microsoft Entra ID)", level=2)
    add_simple_paragraph(doc, "Register a new application in Entra ID (formerly Azure Active Directory) to allow API access:")
    add_simple_bullet(doc, "Navigate to Microsoft Entra ID > App registrations > Click '+ New registration'.")
    add_simple_bullet(doc, "Name it 'security-audit-app' and select 'Accounts in this organizational directory only'.")
    add_simple_bullet(doc, "Copy and save the Application (client) ID and Directory (tenant) ID from the App details page.")
    add_simple_image(doc, "azure-app-registration.jpg", "Azure App Registration details screen showing IDs")
    
    add_simple_heading(doc, "3.2 Generate Client Secret", level=2)
    add_simple_paragraph(doc, "Generate a client secret key to authenticate the App Registration:")
    add_simple_bullet(doc, "Go to Certificates & secrets > Click '+ New client secret'.")
    add_simple_bullet(doc, "Copy the secret 'Value' immediately (it will be hidden after you navigate away).")
    
    add_simple_heading(doc, "3.3 Assign Subscription Reader Role", level=2)
    add_simple_paragraph(doc, "Grant the App Registration access to read subscription resource configurations:")
    add_simple_bullet(doc, "Navigate to Subscriptions > Select your subscription > Access control (IAM).")
    add_simple_bullet(doc, "Click 'Add role assignment' > Search for and select the 'Reader' role.")
    add_simple_bullet(doc, "Assign access to 'User, group, or service principal' and select 'security-audit-app'.")
    add_simple_bullet(doc, "Review and click Save to apply the permissions.")
    add_simple_image(doc, "azure-permissions.jpg", "Assigning Subscription Reader Permissions to App Principal")
    
    add_simple_paragraph(doc, "Deliver the Application ID, Tenant ID, and Client Secret safely to the security assessment team.", bold=True)
    
    doc.save("SOP-GCP-AZURE-PREREQUISITES-CUSTOMER-v1.0.docx")
    print("Generated SOP-GCP-AZURE-PREREQUISITES-CUSTOMER-v1.0.docx")

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 2: INTERNAL PROWLER SCAN EXECUTION & APP DEPLOYMENT
# ═══════════════════════════════════════════════════════════════════════════════
def make_internal_sop():
    doc = Document()
    DOC_ID = "SOP-Prowler-GCP-AZURE-AUDIT-INTERNAL"
    
    add_simple_header(doc, "Cloud Security Assessment: Prowler Audit Execution Guide", DOC_ID)
    
    add_simple_heading(doc, "1. Introduction", level=1)
    add_simple_paragraph(doc, "This internal guide details how to execute security assessment scans using Prowler CLI and how to set up the interactive self-hosted Prowler App using Docker Compose.")
    
    # --- Prowler CLI ---
    add_simple_heading(doc, "2. Running Scans via Prowler CLI", level=1)
    
    add_simple_heading(doc, "2.1 GCP Scan Execution", level=2)
    add_simple_paragraph(doc, "To run the GCP audit, import the customer's shared 'security-audit-key.json' file and run:")
    add_simple_code(doc, 
        "export GOOGLE_APPLICATION_CREDENTIALS=\"/path/to/security-audit-key.json\"\n"
        "export GOOGLE_CLOUD_QUOTA_PROJECT=\"<CUSTOMER_PROJECT_ID>\"\n"
        "prowler gcp"
    )
    
    add_simple_heading(doc, "2.2 Azure Scan Execution", level=2)
    add_simple_paragraph(doc, "To run the Azure audit, configure the customer's tenant, client ID, and secret as environment variables and run Prowler with the Service Principal flag:")
    add_simple_code(doc,
        "export AZURE_CLIENT_ID=\"<CUSTOMER_CLIENT_ID>\"\n"
        "export AZURE_TENANT_ID=\"<CUSTOMER_TENANT_ID>\"\n"
        "export AZURE_CLIENT_SECRET=\"<CUSTOMER_CLIENT_SECRET>\"\n"
        "prowler azure --sp-env-auth"
    )

    # --- Prowler App Docker ---
    add_simple_heading(doc, "3. Setting up the Prowler Web App", level=1)
    add_simple_paragraph(doc, "The Prowler App provides a clean, collaborative dashboard web UI for multi-cloud audits.")
    
    add_simple_heading(doc, "3.1 Download Docker Compose Files", level=2)
    add_simple_paragraph(doc, "Use curl to download the official Prowler Docker configuration files:")
    add_simple_code(doc,
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/docker-compose.yml\n"
        "curl -LO https://raw.githubusercontent.com/prowler-cloud/prowler/refs/heads/master/.env"
    )
    
    add_simple_heading(doc, "3.2 Configure Environment (.env)", level=2)
    add_simple_paragraph(doc, "Open the '.env' file to configure system settings, users, and passwords:")
    add_simple_image(doc, "prowler-env-file.jpg", "Prowler environment config showing port settings and environment options")
    
    add_simple_heading(doc, "3.3 Run the Web App", level=2)
    add_simple_paragraph(doc, "Spin up the Docker containers in detached mode:")
    add_simple_code(doc, "sudo docker compose up -d")
    add_simple_paragraph(doc, "Access the web panel in your browser at http://localhost:3000 (or http://SERVER_IP:3000). Create an admin login upon initial startup.")
    
    # --- Prowler Web App Usage ---
    add_simple_heading(doc, "4. Prowler App Web UI Usage", level=1)
    
    add_simple_heading(doc, "4.1 Connecting Cloud Accounts", level=2)
    add_simple_paragraph(doc, "Log into the Prowler App, navigate to Cloud Providers, select Azure or GCP, and enter the client keys or JSON file details:")
    add_simple_image(doc, "prowler-app-dashboard.jpg", "Prowler App main login/providers page")
    
    add_simple_heading(doc, "4.2 Performing Assessments and Dashboard Views", level=2)
    add_simple_paragraph(doc, "Run the assessment job. The Prowler Dashboard will populate showing Findings by Severity and findings by provider:")
    add_simple_image(doc, "prowler-dashboard-web.jpg", "Interactive Prowler App Dashboard UI")
    
    add_simple_heading(doc, "4.3 Compliance Mappings", level=2)
    add_simple_paragraph(doc, "Navigate to Compliance > Select CIS Benchmark, NIST, or ISO 27001 to view compliance ratios:")
    add_simple_image(doc, "compliance-dashboard.jpg", "Compliance dashboard detailing CIS standard posture ratios")
    
    add_simple_heading(doc, "4.4 Reviewing Misconfigurations", level=2)
    add_simple_paragraph(doc, "Review the active failures listing directly from the UI and download reports:")
    add_simple_image(doc, "azure-misconfigurations.jpg", "Misconfigurations page showing status, severity, and mitigation instructions")
    
    doc.save("SOP-GCP-AZURE-AUDIT-INTERNAL-v1.0.docx")
    print("Generated SOP-GCP-AZURE-AUDIT-INTERNAL-v1.0.docx")

if __name__ == "__main__":
    make_customer_sop()
    make_internal_sop()
